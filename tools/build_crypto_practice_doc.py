from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path(r"C:\Users\angel\flutter\testproject\Crypto_Labs_Scenario_Practice_and_Summary.docx")

# compact_reference_guide preset tokens
PAGE_WIDTH = Inches(8.5)
PAGE_HEIGHT = Inches(11)
MARGIN = Inches(1)
HEADER_FOOTER_DISTANCE = Inches(0.492)
CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
CELL_MARGIN_TOP_BOTTOM = 80
CELL_MARGIN_START_END = 120

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
NAVY = "0B2545"
INK = "25313D"
MUTED = "59636E"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
CALLOUT = "F4F6F9"
PALE_GOLD = "FFF4D6"
GOLD = "7A5A00"
RED = "9B1C1C"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=CELL_MARGIN_TOP_BOTTOM, start=CELL_MARGIN_START_END,
                     bottom=CELL_MARGIN_TOP_BOTTOM, end=CELL_MARGIN_START_END):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_geometry(table, widths_dxa):
    total = sum(widths_dxa)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths_dxa[min(idx, len(widths_dxa) - 1)]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_run_font(run, name="Calibri", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def style_paragraph(paragraph, after=6, before=0, line=1.25, keep_with_next=False):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    fmt.keep_with_next = keep_with_next


def add_page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    set_run_font(run, size=9, color=MUTED)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr)
    run._r.append(fld_char2)


def add_numbering_definition(doc, num_format, text, left=540, hanging=270):
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    abstract_id = max(abstract_ids, default=-1) + 1
    num_ids = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    lvl.append(start)
    fmt = OxmlElement("w:numFmt")
    fmt.set(qn("w:val"), num_format)
    lvl.append(fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), text)
    lvl.append(lvl_text)
    suff = OxmlElement("w:suff")
    suff.set(qn("w:val"), "tab")
    lvl.append(suff)
    ppr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), str(left))
    tabs.append(tab)
    ppr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), str(left))
    ind.set(qn("w:hanging"), str(hanging))
    ppr.append(ind)
    lvl.append(ppr)
    if num_format == "bullet":
        rpr = OxmlElement("w:rPr")
        rfonts = OxmlElement("w:rFonts")
        rfonts.set(qn("w:ascii"), "Calibri")
        rfonts.set(qn("w:hAnsi"), "Calibri")
        rpr.append(rfonts)
        lvl.append(rpr)
    abstract.append(lvl)
    # OOXML requires abstractNum definitions before concrete num instances.
    first_num_index = next(
        (idx for idx, child in enumerate(numbering) if child.tag == qn("w:num")),
        len(numbering),
    )
    numbering.insert(first_num_index, abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def apply_num(paragraph, num_id):
    ppr = paragraph._p.get_or_add_pPr()
    num_pr = ppr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        ppr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num)
    style_paragraph(paragraph, after=4, line=1.25)


def add_heading(doc, text, level=1):
    paragraph = doc.add_paragraph(text, style=f"Heading {level}")
    return paragraph


def add_body(doc, text="", bold_lead=None, italic=False, after=6):
    paragraph = doc.add_paragraph()
    style_paragraph(paragraph, after=after)
    if bold_lead and text.startswith(bold_lead):
        first = paragraph.add_run(bold_lead)
        set_run_font(first, bold=True, color=INK)
        second = paragraph.add_run(text[len(bold_lead):])
        set_run_font(second, italic=italic, color=INK)
    else:
        run = paragraph.add_run(text)
        set_run_font(run, italic=italic, color=INK)
    return paragraph


def add_bullets(doc, items):
    num_id = add_numbering_definition(doc, "bullet", "•")
    for item in items:
        paragraph = doc.add_paragraph()
        apply_num(paragraph, num_id)
        run = paragraph.add_run(item)
        set_run_font(run, color=INK)


def add_numbered(doc, items):
    num_id = add_numbering_definition(doc, "decimal", "%1.")
    for item in items:
        paragraph = doc.add_paragraph()
        apply_num(paragraph, num_id)
        run = paragraph.add_run(item)
        set_run_font(run, color=INK)


def add_code(doc, code):
    paragraph = doc.add_paragraph()
    paragraph.style = doc.styles["Code Block"]
    for idx, line in enumerate(code.strip().splitlines()):
        if idx:
            paragraph.add_run().add_break()
        run = paragraph.add_run(line)
        set_run_font(run, name="Consolas", size=9, color=NAVY)
    return paragraph


def add_callout(doc, label, text, fill=CALLOUT, label_color=DARK_BLUE):
    paragraph = doc.add_paragraph()
    style_paragraph(paragraph, before=4, after=8, line=1.20)
    ppr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    ppr.append(shd)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "160")
    ind.set(qn("w:right"), "160")
    ppr.append(ind)
    r1 = paragraph.add_run(f"{label}: ")
    set_run_font(r1, bold=True, color=label_color)
    r2 = paragraph.add_run(text)
    set_run_font(r2, color=INK)
    return paragraph


def add_two_col_table(doc, rows, widths=(2700, 6660), header=None):
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    if header:
        cells = table.add_row().cells
        for idx, value in enumerate(header):
            cells[idx].text = value
            set_cell_shading(cells[idx], LIGHT_BLUE)
            for run in cells[idx].paragraphs[0].runs:
                set_run_font(run, bold=True, color=NAVY)
        set_repeat_table_header(table.rows[0])
    for left, right in rows:
        cells = table.add_row().cells
        cells[0].text = left
        cells[1].text = right
        for run in cells[0].paragraphs[0].runs:
            set_run_font(run, bold=True, color=DARK_BLUE)
        for run in cells[1].paragraphs[0].runs:
            set_run_font(run, color=INK)
        for cell in cells:
            style_paragraph(cell.paragraphs[0], after=0, line=1.12)
    set_table_geometry(table, list(widths))
    for row_index, row in enumerate(table.rows):
        keep = row_index < len(table.rows) - 1
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.keep_with_next = keep
    spacer = doc.add_paragraph()
    style_paragraph(spacer, after=2)
    return table


def add_mode_table(doc):
    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    headers = ["Mode", "Pattern leakage", "Padding", "One damaged ciphertext byte", "IV/nonce"]
    for idx, value in enumerate(headers):
        cell = table.rows[0].cells[idx]
        cell.text = value
        set_cell_shading(cell, LIGHT_BLUE)
        for run in cell.paragraphs[0].runs:
            set_run_font(run, size=8.5, bold=True, color=NAVY)
        style_paragraph(cell.paragraphs[0], after=0, line=1.05)
    set_repeat_table_header(table.rows[0])
    rows = [
        ("ECB", "Yes", "Yes", "Current block is scrambled", "None"),
        ("CBC", "No", "Yes", "Current block scrambled; same bit flips in next block", "Unique/unpredictable IV"),
        ("CFB", "No", "No", "Same byte changes; next segment/block is scrambled", "Unique IV"),
        ("OFB", "No", "No", "Only the corresponding plaintext byte changes", "Never reuse IV with a key"),
        ("CTR", "No", "No", "Only the corresponding plaintext byte changes", "Never reuse nonce/counter with a key"),
    ]
    for values in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(values):
            cells[idx].text = value
            for run in cells[idx].paragraphs[0].runs:
                set_run_font(run, size=8.5, bold=(idx == 0), color=(DARK_BLUE if idx == 0 else INK))
            style_paragraph(cells[idx].paragraphs[0], after=0, line=1.05)
    set_table_geometry(table, [900, 1440, 900, 3960, 2160])
    for row_index, row in enumerate(table.rows):
        keep = row_index < len(table.rows) - 1
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.keep_with_next = keep
    doc.add_paragraph()


def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def add_scenario(doc, number, title, lab, scenario, objectives, starting_info, tasks, evidence, challenge):
    add_heading(doc, f"Scenario {number}: {title}", 1)
    source_callout = add_callout(doc, "Source coverage", lab, fill=LIGHT_BLUE)
    source_callout.paragraph_format.keep_with_next = True
    add_body(doc, scenario)
    add_heading(doc, "Learning targets", 2)
    add_bullets(doc, objectives)
    add_heading(doc, "Starting information", 2)
    add_two_col_table(doc, starting_info)
    add_heading(doc, "Your tasks", 2)
    add_numbered(doc, tasks)
    add_heading(doc, "Evidence to record", 2)
    add_bullets(doc, evidence)
    add_callout(doc, "Stretch question", challenge, fill=PALE_GOLD, label_color=GOLD)


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    heading_tokens = {
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 14, 7),
        "Heading 3": (12, DARK_BLUE, 10, 5),
    }
    for style_name, (size, color, before, after) in heading_tokens.items():
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    if "Code Block" not in styles:
        code = styles.add_style("Code Block", 1)
    else:
        code = styles["Code Block"]
    code.font.name = "Consolas"
    code._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
    code._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
    code.font.size = Pt(9)
    code.font.color.rgb = RGBColor.from_string(NAVY)
    code.paragraph_format.left_indent = Inches(0.18)
    code.paragraph_format.right_indent = Inches(0.18)
    code.paragraph_format.space_before = Pt(3)
    code.paragraph_format.space_after = Pt(8)
    code.paragraph_format.line_spacing = 1.05
    ppr = code._element.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), LIGHT_GRAY)
    ppr.append(shd)


def configure_sections(doc):
    for section in doc.sections:
        section.page_width = PAGE_WIDTH
        section.page_height = PAGE_HEIGHT
        section.top_margin = MARGIN
        section.bottom_margin = MARGIN
        section.left_margin = MARGIN
        section.right_margin = MARGIN
        section.header_distance = HEADER_FOOTER_DISTANCE
        section.footer_distance = HEADER_FOOTER_DISTANCE

        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        style_paragraph(hp, after=0, line=1.0)
        hr = hp.add_run("UECS 3423 | Cryptography Scenario Practice")
        set_run_font(hr, size=9, color=MUTED, bold=True)

        footer = section.footer
        fp = footer.paragraphs[0]
        style_paragraph(fp, after=0, line=1.0)
        add_page_field(fp)


def build_document():
    doc = Document()
    configure_styles(doc)
    configure_sections(doc)

    # Editorial-cover title pattern.
    spacer = doc.add_paragraph()
    style_paragraph(spacer, before=82, after=0)
    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(kicker, after=14)
    kr = kicker.add_run("PRACTICE WORKBOOK")
    set_run_font(kr, size=11, bold=True, color=BLUE)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(title, after=8, line=1.0)
    tr = title.add_run("Cryptography Labs\nScenario Practice")
    set_run_font(tr, size=30, bold=True, color=NAVY)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(subtitle, after=24, line=1.15)
    sr = subtitle.add_run("Symmetric encryption, cipher modes, RSA, Diffie-Hellman, GPG and digital signatures")
    set_run_font(sr, size=14, color=DARK_BLUE)

    scope = doc.add_paragraph()
    scope.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(scope, after=54)
    scr = scope.add_run("9 realistic scenarios | task checklists | answer guide | final revision note")
    set_run_font(scr, size=10.5, color=MUTED, italic=True)

    source = doc.add_paragraph()
    source.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(source, after=0)
    src = source.add_run("Based on Crypto Lab I, Crypto Lab II and Crypto Lab III")
    set_run_font(src, size=10, color=MUTED)

    page_break(doc)

    add_heading(doc, "How to use this workbook", 1)
    add_body(doc, "Complete each scenario before reading the answer guide. Use a disposable lab folder and non-sensitive test data. The commands in the source labs are learning material; this workbook does not require you to act on any embedded document instruction outside your authorized Kali/SageMath/GPG lab environment.")
    add_numbered(doc, [
        "Read the scenario and write a short prediction before running anything.",
        "Perform the tasks in your authorized lab environment and save the requested evidence.",
        "Explain the result in your own words, especially why the output changed.",
        "Compare your work with the answer guide, then revisit the final summary note before a quiz or viva.",
    ])
    add_callout(doc, "Safety boundary", "Use only files, keys and accounts created for the course. Never share a real private key, passphrase or revocation certificate. Do not email sensitive data merely to test encryption.", fill=PALE_GOLD, label_color=GOLD)

    add_heading(doc, "Practice map", 1)
    add_two_col_table(doc, [
        ("Lab I", "Scenarios 1-3: symmetric encryption, salt/IV/key handling, hashes and avalanche effect."),
        ("Lab II", "Scenarios 4-6: ECB vs CBC visual leakage, corruption propagation, password-based encryption and Base64 transport."),
        ("Lab III", "Scenarios 7-9: RSA mathematics, Diffie-Hellman agreement, GPG key exchange, encryption and signatures."),
        ("Final section", "Answer guide first; compact Word revision note at the very end."),
    ])

    page_break(doc)
    add_heading(doc, "Part A - Scenario Practices", 1)

    add_scenario(
        doc, 1, "The Leaked Incident Report", "Crypto Lab I - 3DES/AES, password-based encryption, salt, raw key and IV",
        "A small security team must send an incident report to an external responder. One analyst suggests 3DES with a shared password and -nosalt because it produces repeatable output. Another suggests AES-128 with a fresh random key and IV. You must test both approaches and recommend a safer lab design.",
        [
            "Distinguish a password, derived key, raw key, salt and IV.",
            "Compare 3DES with AES-128/192/256 at a conceptual level.",
            "Verify successful encryption and decryption without exposing secrets.",
        ],
        [
            ("Plaintext", "incident.txt containing a short fictional incident summary"),
            ("Comparison", "Password-based 3DES versus AES-128 with a 128-bit random raw key and 128-bit IV"),
            ("Inspection tools", "xxd -ps, cmp or diff, and OpenSSL enc"),
            ("Constraint", "No real report, credentials or personal information"),
        ],
        [
            "Encrypt and decrypt the fictional report with password-based 3DES. Confirm that the recovered file is byte-for-byte identical to the original.",
            "Repeat the password-based encryption twice with salt enabled. Compare the ciphertexts and explain why they differ despite using the same password.",
            "Generate 16 random bytes for an AES-128 key and another 16 random bytes for the IV. Encrypt and decrypt using -K and -iv, then verify the recovered file.",
            "Explain the difference among AES-128, AES-192 and AES-256 in key length and number of rounds. State whether AES always has a 128-bit block size.",
            "Write a recommendation that rejects -nosalt for password-based real-world storage and explains why the key and passphrase must not appear in screenshots or command history.",
        ],
        [
            "The two salted ciphertext hashes or hex prefixes, with secrets removed.",
            "A successful byte-for-byte comparison of original and recovered plaintext.",
            "A five-sentence recommendation covering algorithm, salt, IV and secret handling.",
        ],
        "If the IV is not secret, why must it still be generated and handled correctly? Describe what can go wrong when an IV or nonce is reused with the same key.",
    )

    add_scenario(
        doc, 2, "The Suspicious Software Download", "Crypto Lab I - SHA-1, SHA-256, SHA-512 and change detection",
        "A lecturer publishes a lab archive and a SHA-256 digest. A student downloads the archive from a mirror and wants to know whether the file changed. The student also calculates SHA-1 and assumes that a longer digest always means encryption is stronger.",
        [
            "Use cryptographic hashes for integrity comparison.",
            "Separate hashing from encryption and encoding.",
            "Observe the avalanche effect in digest output after a tiny input change.",
        ],
        [
            ("Files", "original.txt and modified.txt differing by one character"),
            ("Digests", "SHA-1, SHA-256 and SHA-512"),
            ("Decision", "Whether a matching digest is sufficient when the published digest source may be untrusted"),
        ],
        [
            "Calculate SHA-1, SHA-256 and SHA-512 for original.txt. Record the digest length in bits and hexadecimal characters.",
            "Change one character, calculate the digests again and compare the outputs. Estimate how much of each digest appears different.",
            "Explain why a hash cannot be decrypted and why it does not provide confidentiality.",
            "Decide whether to trust a matching SHA-256 value obtained from the same compromised mirror as the file. Propose a better verification source.",
            "Explain why SHA-1 should not be selected for new collision-sensitive integrity designs even though it still detects ordinary accidental changes in this exercise.",
        ],
        [
            "Before-and-after digests for all three algorithms.",
            "A table of digest sizes: SHA-1 160 bits/40 hex, SHA-256 256 bits/64 hex, SHA-512 512 bits/128 hex.",
            "A short trust analysis for the source of the reference digest.",
        ],
        "A digest matches. Does that prove who created the file? Identify the additional mechanism needed for origin authentication.",
    )

    add_scenario(
        doc, 3, "One Bit, Three Experiments", "Crypto Lab I - avalanche effect in key, plaintext and ciphertext; CBC vs ECB",
        "A forensic analyst wants to show management why tiny changes can produce large cryptographic effects. You will run three controlled tests: flip one key bit, flip one plaintext bit, and corrupt one ciphertext bit.",
        [
            "Predict avalanche behavior before observing it.",
            "Explain how CBC chaining spreads changes across blocks.",
            "Distinguish encryption avalanche from decryption error propagation.",
        ],
        [
            ("Cipher", "AES-128-CBC, plus AES-128-ECB for comparison"),
            ("Block size", "16 bytes"),
            ("Controls", "Keep every variable fixed except the one bit intentionally changed"),
            ("Comparison", "xxd -ps -c 16, cmp -l and a written block-by-block prediction"),
        ],
        [
            "Encrypt the same plaintext twice, changing one bit of the key only. Predict and observe which ciphertext blocks change.",
            "Restore the original key. Change one bit in plaintext block 1, then encrypt in CBC and ECB. Compare the location and spread of ciphertext changes.",
            "Restore the original plaintext. Corrupt one bit of ciphertext block C1, decrypt in CBC, and identify the affected parts of P1 and P2.",
            "Explain why a corrupted C1 makes P1 unpredictable but flips only the corresponding bit in P2 during CBC decryption.",
            "Write a conclusion separating the cipher's avalanche property from the mode's error-propagation rule.",
        ],
        [
            "Block-aligned hex output for each comparison.",
            "A diagram or paragraph identifying changed blocks.",
            "A final prediction-versus-observation note.",
        ],
        "Why does the experiment need the same key and IV between comparisons? What conclusion becomes invalid if multiple variables change together?",
    )

    add_scenario(
        doc, 4, "The Encrypted Company Logo", "Crypto Lab II - ECB vs CBC bitmap pattern leakage",
        "A company encrypts a bitmap logo before sending it to a printer. The file opens only after its 54-byte BMP header is restored. The ECB version still reveals the logo's outline, while the CBC version looks like noise.",
        [
            "Explain why ECB leaks repeated patterns.",
            "Understand why an encrypted BMP needs a valid header to remain viewable.",
            "Compare confidentiality properties without confusing file format metadata with ciphertext.",
        ],
        [
            ("Input", "A course-provided bitmap such as art.bmp"),
            ("Modes", "AES-128-ECB and AES-128-CBC"),
            ("Header", "First 54 bytes copied from the original BMP into each encrypted output"),
            ("Observation", "Visual comparison only; do not use a personal image"),
        ],
        [
            "Predict which encrypted image will retain visible structure and explain your prediction before running the test.",
            "Encrypt the same bitmap in ECB and CBC. Restore only the 54-byte BMP header so an image viewer can parse the files.",
            "Confirm that the output file size was not accidentally increased by inserting rather than overwriting bytes.",
            "Compare the images and explain how identical plaintext blocks map under ECB versus CBC.",
            "Recommend a mode for protecting structured data and explain why confidentiality should normally be paired with integrity/authentication.",
        ],
        [
            "The original, ECB and CBC images or screenshots.",
            "A 54-byte header check using a hex view.",
            "A paragraph explaining the penguin/logo-style pattern leakage effect.",
        ],
        "Would changing from AES-128-ECB to AES-256-ECB remove the visible pattern? Explain why key size and mode solve different problems.",
    )

    add_scenario(
        doc, 5, "A Damaged Sensor Packet", "Crypto Lab II - corrupted ciphertext in ECB, CBC, CFB, OFB and CTR",
        "An industrial sensor sends a 40-byte encrypted status message. One ciphertext byte is damaged during transmission. The receiver must understand how much plaintext will be lost under each AES mode.",
        [
            "Map one ciphertext error to affected plaintext bytes and blocks.",
            "Explain padding differences between block-like and stream-like modes.",
            "Use controlled comparison rather than relying only on readable text.",
        ],
        [
            ("Plaintext", "Exactly 40 bytes, spanning three AES blocks"),
            ("Modes", "ECB, CBC, full-block CFB, OFB and CTR"),
            ("Damage", "Overwrite ciphertext byte 4 without changing file length"),
            ("Controls", "Same key; same suitable IV where required; -nosalt for byte-position clarity"),
        ],
        [
            "Predict ciphertext file sizes. Explain why ECB/CBC add padding while CFB/OFB/CTR do not for this message.",
            "Encrypt once in each mode, corrupt byte 4, and do not overwrite the damaged files by encrypting again.",
            "Decrypt with the correct key and IV. Use cmp -l and xxd -g 1 -c 16 to identify affected byte positions and blocks.",
            "Complete a results matrix for all five modes: affected block(s), approximate number of changed bytes, and intact blocks.",
            "Select the modes that confine this transmission error to one plaintext byte. Explain the keystream property that prevents propagation.",
        ],
        [
            "Pre-damage and post-damage hex check showing byte 4 changed.",
            "A completed mode comparison matrix.",
            "Block-aligned dumps for at least CBC and CFB.",
        ],
        "Error confinement is not authentication. Explain how an attacker could deliberately modify CTR ciphertext and predictably alter plaintext unless a MAC or authenticated-encryption mode is used.",
    )

    add_scenario(
        doc, 6, "The Email-Friendly Ciphertext", "Crypto Lab II - salt header, IV length, Base64 and password-based key derivation",
        "A team copies encrypted binary into email and finds that the message is corrupted. They switch to Base64 and believe the data is now more secure. Meanwhile, repeated password-based encryption gives different output and an analyst thinks something is wrong.",
        [
            "Explain Base64 as transport encoding, not encryption.",
            "Recognize the OpenSSL Salted__ header and its purpose.",
            "Validate AES IV size and interpret short-IV warnings.",
        ],
        [
            ("Transport", "Email body that can safely carry printable ASCII"),
            ("Password test", "Encrypt the same file twice with a fresh salt each time"),
            ("Raw-key test", "AES-128-CBC with a 16-byte key and deliberately short 8-byte IV"),
        ],
        [
            "Base64-encode a ciphertext and decode it back. Verify that the decoded bytes exactly match the original ciphertext.",
            "Explain why Base64 increases size and provides no confidentiality, integrity or authentication by itself.",
            "Inspect a password-encrypted OpenSSL file. Identify the ASCII Salted__ marker and the following random salt bytes.",
            "Encrypt the same plaintext twice with the same password and salt enabled. Explain why the ciphertexts should differ.",
            "Run the short-IV learning test, record OpenSSL's warning and the zero-padded IV, then state why production code should supply the full required IV/nonce instead of relying on padding behavior.",
        ],
        [
            "A ciphertext-to-Base64-to-ciphertext round-trip comparison.",
            "An annotated 16-byte Salted__ header.",
            "A short explanation of PBKDF, salt and IV roles.",
        ],
        "Why should a modern password-based workflow use a deliberately slow password KDF rather than treating a human password as a raw AES key?",
    )

    add_scenario(
        doc, 7, "RSA Recovery Drill", "Crypto Lab III - RSA key generation and textbook encryption in SageMath",
        "A training system generates two primes and produces an RSA public/private key pair. During a code review, you must verify the mathematics and identify which parts are educational simplifications rather than safe production encryption.",
        [
            "Connect p, q, n, phi(n), e and d.",
            "Verify encryption/decryption using modular exponentiation.",
            "Recognize the security role of padding and adequate key sizes.",
        ],
        [
            ("Primes", "Two independently generated 512-bit primes for the course exercise"),
            ("Public exponent", "e = 65537, provided gcd(e, phi(n)) = 1"),
            ("Message", "A random 128-bit integer m with m < n"),
            ("Environment", "SageMath or SageCell with non-sensitive values"),
        ],
        [
            "Generate p and q, confirm they are distinct primes, then calculate n and phi(n) = (p - 1)(q - 1).",
            "Check gcd(e, phi(n)) = 1 and compute d as the modular inverse of e modulo phi(n).",
            "Encrypt c = m^e mod n, decrypt m' = c^d mod n, and verify m' = m.",
            "Explain which key components may be public and which values must remain secret.",
            "Identify two production concerns in this textbook exercise: the small approximately 1024-bit modulus and the lack of randomized RSA padding such as OAEP.",
        ],
        [
            "Bit lengths of p, q and n.",
            "The gcd check and m == recovered_m result.",
            "A key-component classification and production caveat paragraph.",
        ],
        "If p and q are accidentally equal or predictable, how does that weaken RSA? Explain why prime generation quality matters as much as the equations.",
    )

    add_scenario(
        doc, 8, "The Unauthenticated Key Exchange", "Crypto Lab III - Diffie-Hellman shared secret and man-in-the-middle risk",
        "Alice and Bob calculate matching Diffie-Hellman shared secrets over an untrusted network. They conclude that the connection is secure, but they never authenticate the public values they received.",
        [
            "Compute both sides of a Diffie-Hellman exchange.",
            "Explain why matching mathematics alone does not authenticate a peer.",
            "Separate a shared secret from a usable symmetric session key.",
        ],
        [
            ("Public parameters", "The course-provided 2048-bit prime p and generator g = 5"),
            ("Private values", "Random a and b of at least 128 bits for the exercise"),
            ("Public values", "A = g^a mod p and B = g^b mod p"),
            ("Shared values", "Alice computes B^a mod p; Bob computes A^b mod p"),
        ],
        [
            "Parse p from hexadecimal, verify its bit length and run the requested primality check.",
            "Generate private values a and b, calculate A and B, then compute both shared-secret expressions.",
            "Use modular exponent rules to explain why the two honest-party shared secrets match.",
            "Draw a man-in-the-middle exchange in which Mallory substitutes two public values and creates a separate secret with each victim.",
            "Propose an authenticated design: sign the ephemeral public values or use an authenticated protocol, then process the shared secret through a KDF before using it as an encryption key.",
        ],
        [
            "p bit length, public values and a boolean shared-secret match result; omit private values from shared screenshots.",
            "A three-party man-in-the-middle diagram.",
            "A paragraph distinguishing agreement, authentication and key derivation.",
        ],
        "Why is it unsafe to take the decimal text of the shared secret, truncate it, and use it directly as an AES key?",
    )

    add_scenario(
        doc, 9, "Secure Group Handover with GPG", "Crypto Lab III - GPG key generation, revocation, fingerprints, encryption and digital signatures",
        "Two project groups must exchange a confidential handover note and prove who sent it. One student imports a public key received by email and immediately trusts it. Another stores a revocation certificate in the same shared folder as the public key.",
        [
            "Manage the GPG key lifecycle safely in a lab context.",
            "Verify a public key fingerprint out of band before trusting it.",
            "Differentiate encryption, file signing and key certification.",
        ],
        [
            ("Identities", "Course-only names and addresses; no personal high-value identity"),
            ("Files", "plain.txt, exported armored public keys, encrypted file and signed file"),
            ("Trust check", "Compare the full fingerprint through a separate channel"),
            ("Recovery", "Generate and protect a revocation certificate"),
        ],
        [
            "Generate a course key pair using the settings required by the lab environment and protect the private key with a strong passphrase.",
            "Create a revocation certificate immediately and store it separately with restricted access. Explain why publishing it would invalidate the key.",
            "Export the public key in armored form, import your partner's key, and verify the full fingerprint through an independent channel before certifying it.",
            "Encrypt a fictional handover note to the partner and have the partner decrypt it. Identify which private key is required for decryption.",
            "Clear-sign a message, verify a valid signature, change one character in a copy, and verify again. Record the difference between a good and bad signature result.",
            "Explain the difference among signing a file, signing another person's public key, and signing-then-encrypting a file.",
        ],
        [
            "Public-key fingerprints and verification method; never include passphrases or private-key material.",
            "Successful decryption and good-signature messages.",
            "A failed verification after tampering and a short explanation of why it failed.",
        ],
        "Encryption to a recipient protects confidentiality, while a signature protects origin and integrity. What must you do when both properties are required?",
    )

    page_break(doc)
    add_heading(doc, "Part B - Answer and Debrief Guide", 1)
    add_callout(doc, "Use this section after attempting the scenarios", "The expected results below describe the concepts. Exact ciphertext, keys, salts, signatures and randomly damaged block bytes will differ between runs.", fill=LIGHT_BLUE)

    answer_sections = [
        ("Scenario 1 - Expected reasoning", [
            "Salt is public random input to password-based key derivation; it makes equal passwords produce different derived keys/ciphertexts and frustrates pre-computation. It is not an IV.",
            "An IV initializes a mode. It is usually not secret, but its uniqueness or unpredictability requirements depend on the mode. A raw key is the actual cipher secret; a password should normally pass through a password KDF.",
            "AES-128/192/256 use 128/192/256-bit keys and 10/12/14 rounds. AES block size remains 128 bits in all three.",
            "3DES is included for the course comparison; a new design should use a modern authenticated-encryption construction rather than 3DES or bare AES-CBC without authentication.",
        ]),
        ("Scenario 2 - Expected reasoning", [
            "SHA-1, SHA-256 and SHA-512 output 160, 256 and 512 bits respectively. One tiny input change should make the new digest appear unrelated.",
            "Hashing is one-way integrity processing, not reversible encryption. Base64 is reversible encoding, not hashing or encryption.",
            "A digest from the same compromised location as the file is weak evidence. Obtain the reference digest or signature through a trusted independent channel.",
            "A bare digest does not authenticate the publisher. A verified digital signature can bind integrity evidence to a trusted signing key.",
        ]),
        ("Scenario 3 - Expected reasoning", [
            "Changing one key bit should alter all ciphertext blocks unpredictably because the effective block-cipher mapping changes everywhere.",
            "In ECB, a one-bit plaintext change affects only its own ciphertext block. In CBC, changing P1 changes C1 and, through chaining, every later ciphertext block.",
            "During CBC decryption, corrupting C1 makes P1 unpredictable because D(K, C1) changes strongly; the same corruption mask is XORed into P2, so the corresponding bit(s) in P2 flip. Later blocks are unaffected.",
            "The avalanche effect describes sensitivity inside the cipher; error propagation also depends on how the mode connects blocks.",
        ]),
        ("Scenario 4 - Expected reasoning", [
            "ECB independently maps equal plaintext blocks to equal ciphertext blocks under one key. Repeated pixel regions therefore preserve visible structure.",
            "CBC combines each plaintext block with the previous ciphertext block, so equal plaintext blocks normally encrypt differently. The restored BMP header only makes the file parseable; it does not decrypt the image data.",
            "AES-256-ECB still leaks equal-block patterns. Increasing key length does not repair a weak mode choice.",
            "Confidentiality without integrity permits undetected modification. Prefer an authenticated-encryption design when available.",
        ]),
        ("Scenario 5 - Expected reasoning", [
            "For a 40-byte input, ECB/CBC normally produce 48 ciphertext bytes because of block padding; CFB/OFB/CTR produce 40 bytes.",
            "ECB: the corrupted ciphertext block decrypts unpredictably; other blocks survive. CBC: that block is unpredictable and the same bit difference appears in the next plaintext block.",
            "Full-block CFB: the corresponding plaintext byte changes in the current segment and the next segment/block is disrupted. OFB and CTR: only the corresponding plaintext byte changes because ciphertext is XORed with a keystream that does not depend on earlier ciphertext.",
            "Changed-byte counts in scrambled blocks are typically the full block, but a coincidental matching byte is mathematically possible. Block location is the stronger rule to report.",
        ]),
        ("Scenario 6 - Expected reasoning", [
            "Base64 makes binary safe for text transport and expands the data; anyone can decode it.",
            "OpenSSL password-encrypted output commonly begins with the 8-byte ASCII marker Salted__ followed by salt bytes. Fresh salt explains different ciphertext on repeated runs.",
            "AES has a 16-byte block size, so CBC needs a 16-byte IV. A short-IV warning and zero-padding behavior are teaching observations, not a production design pattern.",
            "A password KDF deliberately adds salt and computational cost so guessing each password is more expensive.",
        ]),
        ("Scenario 7 - Expected reasoning", [
            "n = pq and phi(n) = (p - 1)(q - 1). Choose e coprime to phi(n), then d satisfies ed ≡ 1 mod phi(n). Public key: (n, e); private key: d plus the prime factors and related private parameters.",
            "Encryption is c = m^e mod n; decryption is m = c^d mod n. Correctness depends on RSA's modular arithmetic and m being represented within the modulus domain.",
            "Textbook RSA is deterministic and malleable. Production encryption requires a vetted implementation, adequate keys and randomized padding such as OAEP.",
            "The course's two 512-bit primes yield an approximately 1024-bit modulus and are suitable only for demonstrating the equations, not for protecting real data.",
        ]),
        ("Scenario 8 - Expected reasoning", [
            "A = g^a mod p and B = g^b mod p. Both parties obtain g^(ab) mod p because (g^b)^a = (g^a)^b modulo p.",
            "Plain Diffie-Hellman authenticates nobody. Mallory can replace A and B, creating one shared secret with Alice and another with Bob while relaying traffic.",
            "Authenticate the exchange with signatures/certificates or a proven authenticated key-exchange protocol. Feed the shared secret and transcript context into a KDF to produce correctly sized, separated keys.",
            "Private exponents and derived keys stay secret; public parameters and ephemeral public values may be transmitted but must be validated/authenticated as the protocol requires.",
        ]),
        ("Scenario 9 - Expected reasoning", [
            "A fingerprint must be verified through an independent trusted channel before you rely on the imported key's identity. Importing a key proves only that you possess a copy.",
            "Encrypting to Bob uses Bob's public key; only Bob's corresponding private key can decrypt. Signing uses Alice's private key; anyone with Alice's verified public key can check the signature.",
            "Signing a file authenticates that file. Signing another person's key certifies an identity-key binding. Signing and encrypting combines authenticity/integrity with confidentiality.",
            "A revocation certificate is recovery material. Store it securely and separately because anyone able to publish it can mark the key revoked.",
        ]),
    ]
    for heading, bullets in answer_sections:
        add_heading(doc, heading, 2)
        add_bullets(doc, bullets)

    page_break(doc)
    add_heading(doc, "Part C - Final Summary Note", 1)
    add_callout(doc, "Fast revision method", "Memorize the security goal first: confidentiality, integrity, authentication, non-repudiation support, or key agreement. Then choose the primitive and check its key/IV/nonce/trust requirements.", fill=LIGHT_BLUE)

    add_heading(doc, "1. Core distinctions", 2)
    add_two_col_table(doc, [
        ("Encryption", "Reversible with the correct key; protects confidentiality."),
        ("Hash", "Fixed-length one-way digest; detects change but does not hide data or prove identity by itself."),
        ("Digital signature", "Signer uses a private key; verifier uses the public key; supports integrity and origin authentication."),
        ("Encoding", "Changes representation for storage/transport. Base64 is reversible and provides no security."),
        ("Key agreement", "Lets parties establish shared keying material. Plain Diffie-Hellman does not authenticate peers."),
    ])

    add_heading(doc, "2. Symmetric encryption essentials", 2)
    add_bullets(doc, [
        "AES always has a 128-bit block size. AES-128/192/256 use different key sizes and 10/12/14 rounds.",
        "Salt belongs to password-based key derivation. It is public and helps defeat pre-computation and repeated-output patterns.",
        "IV/nonce belongs to the mode. It is normally public but must satisfy the mode's uniqueness/unpredictability rule. Reuse can be catastrophic.",
        "A key is secret. Do not paste real keys or passphrases into reports, screenshots, emails or shared command history.",
        "-nosalt is useful only for controlled comparison in these labs; it is not a default for real password encryption.",
        "Confidentiality alone does not detect tampering. Prefer authenticated encryption or combine encryption with a correct authentication mechanism.",
    ])

    add_heading(doc, "3. AES mode comparison", 2)
    add_mode_table(doc)

    add_heading(doc, "4. Hash and avalanche facts", 2)
    add_bullets(doc, [
        "SHA-1 = 160 bits = 40 hex characters; SHA-256 = 256 bits = 64 hex; SHA-512 = 512 bits = 128 hex.",
        "A one-bit input change should make roughly half the output bits change on average in a well-designed cryptographic primitive.",
        "A matching hash supports integrity only if the reference digest came from a trusted source.",
        "Changing a key bit, plaintext bit and ciphertext bit are different experiments. Control all other variables.",
        "In CBC encryption, a changed early plaintext block propagates forward. In CBC decryption, a damaged ciphertext block corrupts its own plaintext block and flips corresponding bit(s) in the next block.",
    ])

    add_heading(doc, "5. Public-key essentials", 2)
    add_two_col_table(doc, [
        ("RSA setup", "n = pq; phi(n) = (p - 1)(q - 1); gcd(e, phi) = 1; d = e^-1 mod phi."),
        ("RSA operations", "Encrypt c = m^e mod n; decrypt m = c^d mod n. Textbook RSA needs secure padding in real systems."),
        ("Diffie-Hellman", "A = g^a mod p; B = g^b mod p; shared secret = B^a = A^b = g^(ab) mod p."),
        ("DH warning", "Agreement is not authentication. Prevent MITM with an authenticated protocol/signatures and derive keys with a KDF."),
        ("Public key", "May be shared, but its ownership must be verified."),
        ("Private key", "Never shared; protect with access controls and a strong passphrase where supported."),
    ])

    add_heading(doc, "6. GPG workflow", 2)
    add_numbered(doc, [
        "Generate the course key pair and protect the private key.",
        "Generate a revocation certificate immediately; store it securely and separately.",
        "Export only the public key for sharing, preferably armored for text transport.",
        "Import the partner's public key and verify the full fingerprint out of band before trusting/certifying it.",
        "Encrypt to the recipient's public key for confidentiality.",
        "Sign with your private key for integrity and origin authentication; verify with the signer's trusted public key.",
        "When both properties are needed, sign and encrypt using a well-defined workflow.",
    ])

    add_heading(doc, "7. Common exam and lab traps", 2)
    add_bullets(doc, [
        "Base64 is not encryption, and hashing is not reversible encryption.",
        "ECB pattern leakage is a mode problem; a larger AES key does not fix it.",
        "The IV is not the key, and the salt is not the IV.",
        "A public key received by email is not automatically authentic; verify its fingerprint independently.",
        "A valid decryption does not prove who sent the ciphertext. A verified signature addresses origin and integrity.",
        "Do not regenerate clean ciphertext after corrupting it in an error-propagation experiment.",
        "Count AES blocks in 16-byte lines. Padding can make ECB/CBC ciphertext longer than plaintext.",
        "Educational parameters and textbook algorithms are for understanding, not protecting real sensitive information.",
    ])

    add_heading(doc, "8. One-minute self-test", 2)
    add_numbered(doc, [
        "Which primitive gives confidentiality? Which gives integrity and origin authentication?",
        "Why do two salted encryptions with the same password differ?",
        "Which AES mode visibly leaks repeated bitmap patterns?",
        "What happens to plaintext when one CBC ciphertext block is corrupted?",
        "Why is unauthenticated Diffie-Hellman vulnerable to a man-in-the-middle attack?",
        "Whose key encrypts a GPG message to Bob, and whose key signs a message from Alice?",
    ])
    answers = add_callout(
        doc,
        "Answers",
        "Encryption; signature. Fresh salt changes the derived key/output. ECB. CBC: the current block is scrambled and matching bit(s) flip in the next. DH lacks public-value authentication. Encrypt with Bob's public key; sign with Alice's private key.",
        fill=PALE_GOLD,
        label_color=GOLD,
    )
    style_paragraph(answers, before=2, after=0, line=1.05)
    for run in answers.runs:
        set_run_font(run, size=9.5)

    doc.core_properties.title = "Cryptography Labs Scenario Practice and Summary"
    doc.core_properties.subject = "UECS 3423 scenario practice for Crypto Labs I-III"
    doc.core_properties.author = "OpenAI Codex"
    doc.core_properties.keywords = "cryptography, AES, hash, RSA, Diffie-Hellman, GPG, practice"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
